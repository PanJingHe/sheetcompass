"""
Table region detection and schema extraction (pure Python, no LLM).

Adapted from SpreadsheetBench's table_detection.py.
Loads xlsx → detects sub-table regions → extracts column schemas + data previews.
"""

import openpyxl
from openpyxl.utils import get_column_letter
from dataclasses import dataclass, field
from typing import List, Dict, Optional, Tuple, Any


# ─────────────────────────── Data Classes ───────────────────────────

@dataclass
class TableRegion:
    sheet_name: str
    start_row: int  # 1-based
    start_col: int  # 1-based
    end_row: int
    end_col: int
    has_header: bool = True
    header_rows: int = 1


@dataclass
class ColumnSchema:
    name: str
    col_index: int  # 1-based
    dtype: str  # 'number', 'text', 'date', 'boolean', 'mixed', 'empty'
    null_rate: float = 0.0
    sample_values: List[str] = field(default_factory=list)


@dataclass
class TableSchema:
    table_id: str  # e.g. "Sheet1_T0"
    region: TableRegion
    num_rows: int  # number of **data** rows (excludes header)
    num_cols: int
    columns: List[ColumnSchema] = field(default_factory=list)
    preview_rows: List[List[str]] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "table_id": self.table_id,
            "sheet_name": self.region.sheet_name,
            "range": _data_range(self),
            "full_range": _full_range(self.region),
            "header_row": _header_row(self),
            "data_start_row": _data_start_row(self),
            "data_columns": _data_columns(self),
            "num_rows": self.num_rows,
            "num_cols": self.num_cols,
            "columns": [
                {
                    "name": c.name,
                    "dtype": c.dtype,
                    "null_rate": round(c.null_rate, 3),
                    "sample_values": c.sample_values,
                }
                for c in self.columns
            ],
        }


def _full_range(region: TableRegion) -> str:
    return (
        f"{get_column_letter(region.start_col)}{region.start_row}:"
        f"{get_column_letter(region.end_col)}{region.end_row}"
    )


def _data_range(schema: TableSchema) -> str:
    start_row = _data_start_row(schema)
    last_row = _last_data_row(schema)
    return (
        f"{get_column_letter(schema.region.start_col)}{start_row}:"
        f"{get_column_letter(schema.region.end_col)}{last_row}"
    )


def _header_row(schema: TableSchema) -> Optional[int]:
    if schema.region.has_header and schema.region.header_rows > 0:
        return schema.region.start_row
    return None


def _data_start_row(schema: TableSchema) -> int:
    start_row = schema.region.start_row
    if schema.region.has_header and schema.num_rows > 0:
        start_row += schema.region.header_rows
    return start_row


def _last_data_row(schema: TableSchema) -> int:
    """1-based index of the last **data** row (inclusive). Derived from data_start_row + num_rows."""
    start_row = _data_start_row(schema)
    return start_row + max(schema.num_rows - 1, 0)


def _data_columns(schema: TableSchema) -> str:
    return (
        f"{get_column_letter(schema.region.start_col)}:"
        f"{get_column_letter(schema.region.end_col)}"
    )


# ──────────────────── Step 1: Load Workbook ─────────────────────────

def load_workbook_data(xlsx_path: str) -> Dict[str, List[List[Any]]]:
    """Load xlsx, unmerge cells, return {sheet_name: 2D grid}."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    result = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        merged_map = {}
        for merged_range in ws.merged_cells.ranges:
            top_left_value = ws.cell(merged_range.min_row, merged_range.min_col).value
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for col in range(merged_range.min_col, merged_range.max_col + 1):
                    merged_map[(row, col)] = top_left_value

        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        grid = []
        for r in range(1, max_row + 1):
            row_data = []
            for c in range(1, max_col + 1):
                if (r, c) in merged_map:
                    row_data.append(merged_map[(r, c)])
                else:
                    row_data.append(ws.cell(r, c).value)
            grid.append(row_data)
        result[sheet_name] = grid

    wb.close()
    return result


# ──────────────── Step 2: Table Region Detection ────────────────────

def _is_nonempty(val: Any) -> bool:
    if val is None:
        return False
    if isinstance(val, str) and val.strip() == "":
        return False
    return True


def _find_contiguous_blocks(
    counts: List[int], gap_tolerance: int = 1
) -> List[Tuple[int, int]]:
    blocks = []
    start = None
    gap = 0
    for i, c in enumerate(counts):
        if c > 0:
            if start is None:
                start = i
            gap = 0
        else:
            if start is not None:
                gap += 1
                if gap > gap_tolerance:
                    blocks.append((start, i - gap))
                    start = None
                    gap = 0
    if start is not None:
        end = len(counts) - 1
        while end >= start and counts[end] == 0:
            end -= 1
        blocks.append((start, end))
    return blocks


def _merge_overlapping_regions(regions: List[TableRegion]) -> List[TableRegion]:
    if len(regions) <= 1:
        return regions
    regions.sort(key=lambda r: (r.start_row, r.start_col))
    merged = [regions[0]]
    for reg in regions[1:]:
        last = merged[-1]
        if (
            reg.sheet_name == last.sheet_name
            and reg.start_row <= last.end_row + 2
            and reg.start_col <= last.end_col + 2
            and reg.end_row >= last.start_row - 2
            and reg.end_col >= last.start_col - 2
        ):
            last.start_row = min(last.start_row, reg.start_row)
            last.start_col = min(last.start_col, reg.start_col)
            last.end_row = max(last.end_row, reg.end_row)
            last.end_col = max(last.end_col, reg.end_col)
        else:
            merged.append(reg)
    return merged


def detect_table_regions(
    sheet_name: str,
    grid: List[List[Any]],
    min_rows: int = 2,
    min_cols: int = 2,
    gap_tolerance: int = 1,
) -> List[TableRegion]:
    if not grid or not grid[0]:
        return []

    num_rows = len(grid)
    num_cols = len(grid[0])
    mask = [[_is_nonempty(grid[r][c]) for c in range(num_cols)] for r in range(num_rows)]
    row_counts = [sum(mask[r]) for r in range(num_rows)]
    col_counts = [sum(mask[r][c] for r in range(num_rows)) for c in range(num_cols)]

    row_blocks = _find_contiguous_blocks(row_counts, gap_tolerance=gap_tolerance)
    col_blocks = _find_contiguous_blocks(col_counts, gap_tolerance=gap_tolerance)

    regions = []
    for r_start, r_end in row_blocks:
        for c_start, c_end in col_blocks:
            filled = 0
            total = (r_end - r_start + 1) * (c_end - c_start + 1)
            for r in range(r_start, r_end + 1):
                for c in range(c_start, c_end + 1):
                    if mask[r][c]:
                        filled += 1
            if filled < total * 0.2:
                continue
            if (r_end - r_start + 1) < min_rows or (c_end - c_start + 1) < min_cols:
                continue

            first_row = grid[r_start]
            text_count = sum(
                1 for c in range(c_start, c_end + 1)
                if isinstance(first_row[c], str) and first_row[c].strip()
            )
            has_header = text_count >= (c_end - c_start + 1) * 0.5
            regions.append(TableRegion(
                sheet_name=sheet_name,
                start_row=r_start + 1,
                start_col=c_start + 1,
                end_row=r_end + 1,
                end_col=c_end + 1,
                has_header=has_header,
                header_rows=1 if has_header else 0,
            ))

    regions = _merge_overlapping_regions(regions)

    if not regions and any(row_counts):
        r_start = next(i for i, c in enumerate(row_counts) if c > 0)
        r_end = num_rows - 1 - next(i for i, c in enumerate(reversed(row_counts)) if c > 0)
        c_start = next(i for i, c in enumerate(col_counts) if c > 0)
        c_end = num_cols - 1 - next(i for i, c in enumerate(reversed(col_counts)) if c > 0)
        regions.append(TableRegion(
            sheet_name=sheet_name,
            start_row=r_start + 1, start_col=c_start + 1,
            end_row=r_end + 1, end_col=c_end + 1,
        ))
    return regions


# ──────────────── Step 3: Schema Extraction ─────────────────────────

def _infer_dtype(values: List[Any]) -> str:
    non_null = [v for v in values if v is not None and str(v).strip() != ""]
    if not non_null:
        return "empty"
    type_counts = {"number": 0, "text": 0, "date": 0, "boolean": 0}
    for v in non_null:
        if isinstance(v, bool):
            type_counts["boolean"] += 1
        elif isinstance(v, (int, float)):
            type_counts["number"] += 1
        elif hasattr(v, "strftime"):
            type_counts["date"] += 1
        elif isinstance(v, str):
            try:
                float(v.replace(",", "").replace("%", ""))
                type_counts["number"] += 1
            except ValueError:
                type_counts["text"] += 1
        else:
            type_counts["text"] += 1
    dominant = max(type_counts, key=type_counts.get)
    if type_counts[dominant] < len(non_null) * 0.7:
        return "mixed"
    return dominant


PREVIEW_ROWS = 3


def extract_schema(
    sheet_name: str,
    grid: List[List[Any]],
    region: TableRegion,
    table_index: int = 0,
    preview_rows: int = PREVIEW_ROWS,
) -> TableSchema:
    r1 = region.start_row - 1
    r2 = region.end_row - 1
    c1 = region.start_col - 1
    c2 = region.end_col - 1

    headers = []
    data_start = r1
    if region.has_header:
        for c in range(c1, c2 + 1):
            val = grid[r1][c]
            headers.append(str(val).strip() if val is not None else f"Col_{get_column_letter(c + 1)}")
        data_start = r1 + region.header_rows
    else:
        for c in range(c1, c2 + 1):
            headers.append(f"Col_{get_column_letter(c + 1)}")

    num_data_rows = r2 - data_start + 1
    num_cols = c2 - c1 + 1

    columns = []
    for ci, header in enumerate(headers):
        col_idx = c1 + ci
        col_values = [
            grid[r][col_idx] if r < len(grid) and col_idx < len(grid[r]) else None
            for r in range(data_start, r2 + 1)
        ]
        non_null = [v for v in col_values if v is not None and str(v).strip() != ""]
        null_rate = 1.0 - (len(non_null) / max(len(col_values), 1))
        dtype = _infer_dtype(col_values)
        samples = []
        for v in non_null[:5]:
            s = str(v)
            if len(s) > 50:
                s = s[:47] + "..."
            samples.append(s)
        columns.append(ColumnSchema(
            name=header, col_index=col_idx + 1,
            dtype=dtype, null_rate=null_rate, sample_values=samples,
        ))

    preview = []
    preview_end = min(data_start + preview_rows, r2 + 1)
    for r in range(data_start, preview_end):
        row_data = []
        for c in range(c1, c2 + 1):
            val = grid[r][c] if r < len(grid) and c < len(grid[r]) else None
            if val is None:
                row_data.append("")
            else:
                s = str(val)
                if len(s) > 50:
                    s = s[:47] + "..."
                row_data.append(s)
        preview.append(row_data)

    return TableSchema(
        table_id=f"{sheet_name}_T{table_index}",
        region=region,
        num_rows=max(num_data_rows, 0),
        num_cols=num_cols,
        columns=columns,
        preview_rows=preview,
    )


# ──────────────── Main Entry ────────────────────────────────────────

def build_tables_preview_for_semantic(
    schemas: List[TableSchema],
    *,
    max_preview_rows: int = 4,
) -> List[dict]:
    """
    Per-table header-keyed sample rows for semantic LLM prompts (column meaning from names + values).
    """
    out: List[dict] = []
    n = max(0, max_preview_rows)
    for schema in schemas:
        headers = [c.name for c in schema.columns]
        keyed_rows: List[dict] = []
        for raw in schema.preview_rows[:n]:
            row = raw if isinstance(raw, list) else []
            keyed: dict = {}
            for i, h in enumerate(headers):
                keyed[h] = row[i] if i < len(row) else ""
            keyed_rows.append(keyed)
        out.append({
            "table_id": schema.table_id,
            "sheet_name": schema.region.sheet_name,
            "num_cols": schema.num_cols,
            "range": _data_range(schema),
            "preview_rows": keyed_rows,
        })
    return out


def process_spreadsheet(xlsx_path: str) -> List[TableSchema]:
    """Load xlsx → detect tables → extract schemas."""
    workbook_data = load_workbook_data(xlsx_path)
    all_schemas = []
    global_idx = 0
    for sheet_name, grid in workbook_data.items():
        regions = detect_table_regions(sheet_name, grid)
        for region in regions:
            schema = extract_schema(sheet_name, grid, region, table_index=global_idx)
            all_schemas.append(schema)
            global_idx += 1
    return all_schemas


def format_sheet_state_sheetcopilot_style(schemas: List[TableSchema]) -> str:
    """
    One line per detected table, SheetCopilot-style column listing plus explicit data bounds.

    States both (a) total rows in the detected block and (b) data row count / row numbers so
    prompts stay consistent with schema ``num_rows`` and ``range(data_start_row, ...)``.
    """
    if not schemas:
        return (
            "Sheet state: No contiguous tables were auto-detected from the workbook. "
            "Rely on the input file path, sheet names in the instruction, and the schema section below."
        )
    lines = []
    for i, s in enumerate(schemas):
        active = " (active)" if i == 0 else ""
        letters_headers = []
        for c in s.columns:
            letter = get_column_letter(c.col_index)
            name = (c.name or "").replace('"', "'")
            letters_headers.append(f'{letter}: "{name}"')
        headers_str = ", ".join(letters_headers) if letters_headers else "(no column headers)"
        total_rows = s.region.end_row - s.region.start_row + 1
        hr = _header_row(s)
        if hr is not None and s.num_rows > 0:
            ds = _data_start_row(s)
            last_data = _last_data_row(s)
            hi = ds + s.num_rows
            row_note = (
                f"header row {hr}; {s.num_rows} data rows on rows {ds}–{last_data} "
                f"(use ``range({ds}, {hi})`` for full data); {total_rows} rows in the detected block including header"
            )
        elif hr is not None:
            row_note = f"header row {hr}; {total_rows} rows in the detected block (no data rows counted)"
        elif s.num_rows > 0:
            ds = _data_start_row(s)
            last_data = _last_data_row(s)
            hi = ds + s.num_rows
            row_note = (
                f"{s.num_rows} data rows on rows {ds}–{last_data} "
                f"(``range({ds}, {hi})``); {total_rows} rows in the detected block"
            )
        else:
            row_note = f"{total_rows} rows in the detected block; no header row detected"
        lines.append(
            f'Sheet "{s.region.sheet_name}"{active} has {len(s.columns)} columns '
            f"(Headers are {headers_str}). {row_note}."
        )
    return "Sheet state:\n" + "\n".join(f"- {ln}" for ln in lines)


def format_raw_schemas(schemas: List[TableSchema]) -> str:
    """Format schemas with data preview as markdown for LLM prompt."""
    parts = []
    for s in schemas:
        lines = [
            f"### Table: {s.table_id}",
            f"- Sheet: {s.region.sheet_name}",
            f"- Data range: {_data_range(s)}",
            f"- Full detected range: {_full_range(s.region)}",
            f"- Header row: {_header_row(s) if _header_row(s) is not None else '(none)'}",
            f"- Data start row: {_data_start_row(s)}",
            f"- num_data_rows: {s.num_rows} (Python: range(data_start_row, data_start_row + num_data_rows))",
            f"- Data columns: {_data_columns(s)}",
            f"- Columns: {s.num_cols}",
            f"- Column stats:",
        ]
        for c in s.columns:
            lines.append(f"  - {c.name} ({c.dtype}, null_rate={c.null_rate:.1%})")
        if s.preview_rows:
            lines.append("")
            lines.append(f"**Data preview** (first {len(s.preview_rows)} of {s.num_rows} rows):")
            lines.append("")
            col_names = [c.name for c in s.columns]
            lines.append("| " + " | ".join(col_names) + " |")
            lines.append("| " + " | ".join(["---"] * len(col_names)) + " |")
            for row in s.preview_rows:
                cells = [str(v).replace("|", "\\|") if v else "" for v in row]
                lines.append("| " + " | ".join(cells) + " |")
        parts.append("\n".join(lines))
    return "\n\n".join(parts)
