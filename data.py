import json
import os

from openpyxl import load_workbook


def read_excel_preview(excel_path, preview_rows=5):
    """Read a compact text preview from the first few rows of each sheet."""
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    wb = load_workbook(excel_path, data_only=True, read_only=True)
    chunks = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=preview_rows, values_only=True):
            rows.append("\t".join("" if cell is None else str(cell) for cell in row))
        chunks.append(f"Sheet Name: {ws.title}\n" + "\n".join(rows))
    wb.close()
    return ("-" * 50 + "\n").join(chunks) if chunks else ""


class SpreadsheetDataLoader:
    """SpreadsheetBench dataset loader."""

    benchmark_name = "spreadsheetbench"

    def __init__(self, data_root, dataset_name="sample_data_200", num_samples=-1, preview_rows=5, sample_id=None):
        self.data_root = data_root
        self.dataset_name = dataset_name
        self.num_samples = num_samples
        self.preview_rows = preview_rows
        self.sample_id = str(sample_id) if sample_id is not None else None
        self.dataset_path = os.path.join(data_root, dataset_name)
        self.dataset_json_path = os.path.join(self.dataset_path, "dataset.json")
        if not os.path.exists(self.dataset_json_path):
            raise FileNotFoundError(f"Dataset file not found: {self.dataset_json_path}")

    def load_dataset(self):
        with open(self.dataset_json_path, "r", encoding="utf-8") as fp:
            dataset = json.load(fp)
        if self.sample_id is not None:
            dataset = [sample for sample in dataset if str(sample["id"]) == self.sample_id]
        return dataset[:self.num_samples] if self.num_samples > 0 else dataset

    def get_excel_path(self, sample):
        sample_id = sample["spreadsheet_path"].lstrip("spreadsheet/")
        file_name = f"1_{sample_id}_input.xlsx"
        return os.path.join(self.dataset_path, sample["spreadsheet_path"], file_name)

    def load_sample(self, sample):
        excel_path = self.get_excel_path(sample)
        return {
            "benchmark": self.benchmark_name,
            "id": str(sample["id"]),
            "instruction": sample["instruction"],
            "spreadsheet_path": sample["spreadsheet_path"],
            "instruction_type": sample["instruction_type"],
            "answer_position": sample["answer_position"],
            "excel_path": excel_path,
            "excel_content": read_excel_preview(excel_path, self.preview_rows),
        }

    def load_all_samples(self):
        samples = []
        for sample in self.load_dataset():
            try:
                samples.append(self.load_sample(sample))
            except FileNotFoundError as e:
                print(f"[Warning] Skipping sample {sample['id']}: {e}")
        return samples


class SheetCopilotDataLoader:
    """SheetCopilot benchmark loader."""

    benchmark_name = "sheetcopilot"

    def __init__(self, data_root, num_samples=-1, preview_rows=5, sample_id=None, gt_version="task_sheet_answers_v2"):
        self.data_root = data_root
        self.num_samples = num_samples
        self.preview_rows = preview_rows
        self.sample_id = str(sample_id) if sample_id is not None else None
        self.gt_version = gt_version
        self.dataset_xlsx_path = os.path.join(data_root, "dataset.xlsx")
        self.task_sheets_dir = os.path.join(data_root, "task_sheets")
        self.gt_root = os.path.join(data_root, gt_version)
        if not os.path.exists(self.dataset_xlsx_path):
            raise FileNotFoundError(f"Dataset file not found: {self.dataset_xlsx_path}")
        if not os.path.isdir(self.task_sheets_dir):
            raise FileNotFoundError(f"Task sheets directory not found: {self.task_sheets_dir}")
        if not os.path.isdir(self.gt_root):
            raise FileNotFoundError(f"Ground-truth directory not found: {self.gt_root}")

    def load_dataset(self):
        wb = load_workbook(self.dataset_xlsx_path, data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        header = list(rows[0])
        header_index = {name: idx for idx, name in enumerate(header)}

        def row_value(row, column_name, default=""):
            idx = header_index.get(column_name)
            if idx is None:
                return default
            value = row[idx]
            return default if value is None else value

        dataset = []
        for row in rows[1:]:
            if not any(value is not None for value in row):
                continue
            sheet_name = str(row_value(row, "Sheet Name"))
            task_no = str(row_value(row, "No."))
            task_id = f"{task_no}_{sheet_name}"
            sample = {
                "id": task_id,
                "task_no": task_no,
                "sheet_name": sheet_name,
                "context": row_value(row, "Context"),
                "instruction": row_value(row, "Instructions"),
                "instruction_type": row_value(row, "Categories"),
                "categories": row_value(row, "Categories"),
                "atomic_actions": row_value(row, "Atomic actions"),
                "source": row_value(row, "Source"),
            }
            dataset.append(sample)

        if self.sample_id is not None:
            dataset = [sample for sample in dataset if str(sample["id"]) == self.sample_id]
        return dataset[:self.num_samples] if self.num_samples > 0 else dataset

    def get_excel_path(self, sample):
        return os.path.join(self.task_sheets_dir, f"{sample['sheet_name']}.xlsx")

    def get_gt_dir(self, sample):
        return os.path.join(self.gt_root, sample["sheet_name"], sample["id"])

    def load_sample(self, sample):
        workbook_path = self.get_excel_path(sample)
        gt_dir = self.get_gt_dir(sample)
        if not os.path.isdir(gt_dir):
            raise FileNotFoundError(f"GT folder not found: {gt_dir}")
        return {
            "benchmark": self.benchmark_name,
            "id": sample["id"],
            "task_no": sample["task_no"],
            "sheet_name": sample["sheet_name"],
            "context": sample["context"],
            "instruction": sample["instruction"],
            "instruction_type": sample["instruction_type"],
            "categories": sample["categories"],
            "atomic_actions": sample["atomic_actions"],
            "source": sample["source"],
            "workbook_path": workbook_path,
            "gt_dir": gt_dir,
            "excel_path": workbook_path,
            "excel_content": read_excel_preview(workbook_path, self.preview_rows),
        }

    def load_all_samples(self):
        samples = []
        for sample in self.load_dataset():
            try:
                samples.append(self.load_sample(sample))
            except FileNotFoundError as e:
                print(f"[Warning] Skipping sample {sample['id']}: {e}")
        return samples
